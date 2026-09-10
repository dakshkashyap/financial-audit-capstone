"""Generate the weekly progress report as a .docx for upload to Google Docs.

Usage:  python scripts/make_weekly_report_docx.py
Output: Weekly_Report_Irvin.docx
"""

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

REPO = "https://github.com/dakshkashyap/financial-audit-capstone"
COMMIT = REPO + "/commit/"


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        docx.opc.constants.RELATIONSHIP_TYPE.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1155CC")
    rPr.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rPr.append(underline)

    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    rPr.append(size)

    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def c(sha):
    """Shorthand for a commit link run."""
    return ("link", COMMIT + sha, sha)


def write_runs(paragraph, runs):
    for item in runs:
        if isinstance(item, tuple) and item[0] == "link":
            add_hyperlink(paragraph, item[1], item[2])
        elif isinstance(item, tuple) and item[0] == "bold":
            r = paragraph.add_run(item[1])
            r.bold = True
            r.font.size = Pt(9)
        else:
            r = paragraph.add_run(item)
            r.font.size = Pt(9)


def fill_cell(cell, bullets):
    cell.text = ""
    for i, bullet in enumerate(bullets):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(0)
        write_runs(p, bullet)


def simple_cell(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)


MONTHS = [
    (
        "Month of May",
        [
            (
                "Week 1",
                "May 15, 2026",
                "Shared my reading of the professor's paper and where I thought the "
                "biggest gap was",
                [
                    [
                        "Read and analyzed ",
                        ("bold", "Automating Financial Statement Audits with Large "
                                 "Language Models"),
                        " (AuditBench, arXiv:2506.17282), the paper provided by the "
                        "professor"
                    ],
                    [
                        "Focused on the weakest reported result — citing the correct "
                        "accounting standard — and argued it is a design problem rather "
                        "than a model problem, since the model is recalling rule numbers "
                        "from memory instead of looking them up"
                    ],
                ],
            ),
            (
                "Week 2",
                "May 22, 2026",
                "Group discussion on dataset collection and the planned meeting with the "
                "RBC Borealis engineer on their financial model ATOM",
                [
                    [
                        "Researched related work in the LLM financial auditing space and "
                        "the data sources available for real filings (SEC EDGAR)"
                    ],
                    ["Set up local environment and repo access"],
                ],
            ),
            (
                "Week 3",
                "May 29, 2026",
                "",
                [
                    [
                        "Studied retrieval-based and knowledge-graph approaches as a way "
                        "to ground citations in an authoritative source instead of model "
                        "memory"
                    ]
                ],
            ),
        ],
    ),
    (
        "Month of June",
        [
            (
                "Week 1",
                "Jun 5, 2026",
                "Met with Mohsen on Jun 4th",
                [
                    [
                        "Reviewed the AuditBench replication the team had produced and "
                        "agreed to take the label-to-concept mapping stage (Stage 1)"
                    ]
                ],
            ),
            (
                "Week 2",
                "Jun 12, 2026",
                "",
                [
                    [
                        "Studied how the baseline pipeline scores citations, and where "
                        "an XBRL concept lookup could replace free-text generation"
                    ]
                ],
            ),
            (
                "Week 3",
                "Jun 19, 2026",
                "",
                [
                    [
                        "Designed the Stage 1 approach: map each statement line-item "
                        "label to its official XBRL concept so downstream citation "
                        "lookup has a grounded key to work from"
                    ]
                ],
            ),
            (
                "Week 4",
                "Jun 26, 2026",
                "Took ownership of Stage 1 (label to XBRL concept mapping)",
                [
                    [
                        "Built the EDGAR Mapper for Stage 1: maps a financial statement "
                        "line-item label to its official XBRL concept, with an evaluation "
                        "harness to score it ",
                        c("d2a1cb3"),
                    ],
                    [
                        "Fixed statement-type inference so error splits without a "
                        "Sheet_type field are classified from content instead ",
                        c("0d34c06"),
                    ],
                    [
                        "Wrote full Stage 1 documentation covering the pipeline, file "
                        "breakdown, eval results, and integration with Manish's stage ",
                        c("dc48d61"),
                    ],
                ],
            ),
        ],
    ),
    (
        "Month of July",
        [
            (
                "Week 1",
                "Jul 3, 2026",
                "Shared my written research analysis with the team",
                [
                    [
                        "Wrote RESEARCH_ANALYSIS.md, a professor-facing comparison of "
                        "three papers against our approach: ",
                        ("bold", "AuditBench"),
                        " (arXiv:2506.17282), ",
                        ("bold", "FinAuditing"),
                        " (arXiv:2510.08886) and ",
                        ("bold", "AuditFlow"),
                        " (arXiv:2606.03031) ",
                        c("7622f73"),
                    ],
                    [
                        "Key takeaway carried into our design: AuditFlow's principle of "
                        "separating LLM-guided search from deterministic verification, "
                        "and FinAuditing's use of real XBRL filings with rule-based "
                        "labels instead of model-written ones"
                    ],
                    [
                        "Documented the gaps in our own architecture alongside the "
                        "comparison rather than only the strengths"
                    ],
                ],
            ),
            (
                "Week 2",
                "Jul 10, 2026",
                "",
                [
                    [
                        "Designed the ",
                        ("bold", "Taxonomy Knowledge Agent"),
                        ": an MCP server wrapping the FASB US-GAAP taxonomy graph so any "
                        "agent in the pipeline can look up accounting rules through tools "
                        "and never generate a citation from memory ",
                        c("8fd7d87"),
                    ],
                    [
                        "Wrote the next-phase architecture plan adapting AuditFlow's "
                        "verification principle — agents share a central MCP server "
                        "holding the verified evidence and cannot bypass the tools ",
                        c("b848a8b"),
                    ],
                    [
                        "Documented how state is stored and shared across MCP tool calls ",
                        c("04d3f71"),
                    ],
                ],
            ),
            (
                "Week 3",
                "Jul 17, 2026",
                "",
                [
                    [
                        "Released EDGAR Mapper v2 with its documentation and dataset "
                        "links ",
                        c("591e5d7"),
                    ],
                    [
                        "Updated the July meeting notes and the taxonomy MCP agent design ",
                        c("e11f3a2"),
                    ],
                ],
            ),
            (
                "Week 4",
                "Jul 24, 2026",
                "",
                [
                    [
                        "Analyzed the FinMR benchmark and how its DQC-rule violations "
                        "provide root-cause ground truth — a reported value and a "
                        "calculated value per violation, rather than a pass/fail label ",
                        c("afa52aa"),
                    ]
                ],
            ),
            (
                "Week 5",
                "Jul 31, 2026",
                "Proposed using an MCP server with an AI agent for the citation stage, "
                "and how to define \"works better\"",
                [
                    [
                        "Built a SEC EDGAR XBRL client to pull real company facts directly "
                        "from the SEC API, no API key required ",
                        c("a0e8de7"),
                    ],
                    [
                        ("bold", "Built the Citation MCP agent: "),
                        "wrapped the TaxonomyGraph in an MCP server exposing "
                        "get_candidates, validate_citation, get_concept_info and "
                        "store_pick, so an AI agent can only choose from real "
                        "taxonomy-grounded citations instead of generating them from "
                        "memory ",
                        c("6a353c5"),
                    ],
                    [
                        "Wrote a Current vs Proposed architecture brief with diagrams for "
                        "the professor, plus a diagram-first walkthrough of the "
                        "TaxonomyGraph and MCP flow ",
                        c("5de40e4"),
                        "  ",
                        c("34b28de"),
                    ],
                    [
                        "Wired up Claude Haiku 4.5 and ran the citation experiment across "
                        "three conditions — lookup only, AI tool-locked, and an oracle "
                        "ceiling — on 42 items. Result: ",
                        ("bold", "0 invented citations"),
                        ", 100% of accepted answers verified against the taxonomy, with a "
                        "written rationale per pick ",
                        c("889922b"),
                    ],
                ],
            ),
        ],
    ),
    (
        "Month of August",
        [
            (
                "Week 1",
                "Aug 7, 2026",
                "Presented the AuditPatch pipeline and the three deliverables: "
                "explainability, clear methodology, root-cause ground truth",
                [
                    [
                        ("bold", "Built the AuditPatch repair pipeline on FinMR: "),
                        "detect the rule violation, localize the single fact that caused "
                        "it, generate a minimal one-value repair, revalidate the whole "
                        "filing, attach a taxonomy-grounded ASC citation, and emit a "
                        "machine-checkable certificate ",
                        c("7c43575"),
                    ],
                    [
                        "Ran it on all ",
                        ("bold", "332 real SEC filings"),
                        " in FinMR: 178 repairs proposed, ",
                        ("bold", "145 exactly matched ground truth (81.5%)"),
                        ", ",
                        ("bold", "0 regressions"),
                        ", 1 value changed per repair, 93% carrying a verified citation",
                    ],
                    [
                        "Wrote a professor brief tying the results to the three "
                        "deliverables ",
                        c("666bf62"),
                        " and a beginner-friendly presentation walkthrough ",
                        c("6ee78aa"),
                    ],
                    [
                        "Documented the LLM explainability experiment in the "
                        "presentation, including the honest finding that tool-locking "
                        "removed hallucination but accuracy stayed capped at the 26.2% "
                        "candidate-list ceiling ",
                        c("ba87b20"),
                    ],
                    [
                        "Added a short architecture deck framing the six-stage pipeline "
                        "as the settled contribution, with each stage's method as the "
                        "part we iterate on ",
                        c("141a697"),
                    ],
                ],
            ),
            (
                "Week 2",
                "Aug 14, 2026",
                "Meeting cancelled due to unavailability. Met with Mohsen as a group on "
                "Aug 13th to share the overall progress and next steps",
                [["Prepared for final exams for other courses"]],
            ),
            (
                "Week 3",
                "Aug 21, 2026",
                "",
                [
                    [
                        "Related-work review of models fine-tuned for audit and XBRL — "
                        "AuditWen, FinLoRA, RKEFino1, FinTag, XBRL-Agent — to position "
                        "our contribution. Finding: existing work reads and tags filings; "
                        "none localize a root cause, propose a minimal repair, or "
                        "revalidate it. XBRL-Agent (ICAIF 2024) is the closest precedent "
                        "for our tool-grounding argument"
                    ],
                    [
                        "Drafted an evaluation plan for the explanation layer: automatic "
                        "faithfulness checks on every generated explanation, a "
                        "reconstruction test as a proxy for usefulness, counterfactual "
                        "sensitivity, and ablations across no-tools / tools / "
                        "tools-plus-validation"
                    ],
                ],
            ),
            ("Week 4", "Aug 28, 2026", "", [[""]]),
        ],
    ),
]


def build():
    doc = docx.Document()

    for section in doc.sections:
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(10)

    title = doc.add_paragraph()
    tr = title.add_run("Weekly Report — Irvin")
    tr.bold = True
    tr.font.size = Pt(16)

    repo_p = doc.add_paragraph()
    repo_p.add_run("Github repo: ").font.size = Pt(10)
    add_hyperlink(repo_p, REPO, REPO)

    for month_name, rows in MONTHS:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(14)
        hr = h.add_run(month_name)
        hr.bold = True
        hr.font.size = Pt(13)

        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        headers = [
            "Week",
            "Meeting Date",
            "Meeting minutes",
            "Description of work done during the week",
        ]
        for i, text in enumerate(headers):
            simple_cell(table.rows[0].cells[i], text, bold=True, size=10)

        widths = [Inches(0.7), Inches(1.0), Inches(2.0), Inches(4.0)]

        for week, date, minutes, bullets in rows:
            cells = table.add_row().cells
            simple_cell(cells[0], week)
            simple_cell(cells[1], date)
            simple_cell(cells[2], minutes)
            fill_cell(cells[3], bullets)

        for row in table.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = w

    doc.add_page_break()
    s = doc.add_paragraph()
    sr = s.add_run("Summary of contributions")
    sr.bold = True
    sr.font.size = Pt(13)

    for heading, body in [
        (
            "Stage 1 — EDGAR Mapper.",
            " Label-to-XBRL-concept mapper with its own evaluation harness and "
            "documentation.",
        ),
        (
            "Citation MCP agent.",
            " Wrapped the FASB US-GAAP taxonomy in an MCP server so an AI agent must "
            "pick from real, verified citations. Measured across three conditions: zero "
            "hallucinated citations, every accepted answer verified, and a written "
            "justification for each pick.",
        ),
        (
            "AuditPatch pipeline.",
            " End-to-end detect, localize, repair, revalidate, cite and certify on 332 "
            "real SEC filings, achieving 81.5% exact repair match with zero regressions.",
        ),
        (
            "Documentation.",
            " Architecture comparison, taxonomy and MCP flow walkthrough, professor "
            "brief, full presentation, and the short architecture deck.",
        ),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.add_run(heading).bold = True
        p.add_run(body)

    out = "Weekly_Report_Irvin.docx"
    doc.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    build()
